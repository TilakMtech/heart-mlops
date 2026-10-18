"""Build a presentation deck for the assignment walkthrough (MS Teams screen share).

Reads live project outputs, so re-run it after retraining or adding screenshots:
  - reports/model_comparison.csv, models/model_meta.json, reports/figures/eda_summary.json
  - reports/figures/*.png, docs/architecture.png, screenshots/*.png
  - presenter name from reports/REPORT.md, repo URL from `git remote get-url origin`

    python scripts/build_deck.py          # -> reports/Heart_Disease_MLOps_Presentation.pptx
"""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "Heart_Disease_MLOps_Presentation.pptx"
FIG = ROOT / "reports" / "figures"
SHOTS = ROOT / "screenshots"

# palette: deep navy dominant, heart-red accent, teal support (matches EDA figures)
NAVY = RGBColor(0x14, 0x21, 0x3D)
NAVY_2 = RGBColor(0x1F, 0x30, 0x55)
RED = RGBColor(0xE6, 0x39, 0x46)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
SAND = RGBColor(0xF6, 0xF4, 0xEF)
INK = RGBColor(0x22, 0x2B, 0x38)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xDC, 0xE1)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)


# --------------------------------------------------------------------------- data
def load_context() -> dict:
    ctx = {"name": "Student", "repo": "https://github.com/TilakMtech/heart-mlops", "tests": "37"}
    report = ROOT / "reports" / "REPORT.md"
    if report.exists():
        m = re.search(r"\*\*Student:\*\*\s*([^*·\n]+)", report.read_text())
        if m:
            ctx["name"] = m.group(1).strip()
    try:
        url = subprocess.run(
            ["git", "remote", "get-url", "origin"], cwd=ROOT, capture_output=True, text=True, check=True
        ).stdout.strip()
        ctx["repo"] = re.sub(r"\.git$", "", url)
    except Exception:
        pass
    try:
        out = subprocess.run(
            ["python", "-m", "pytest", "--collect-only", "-q"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            timeout=120,
        ).stdout
        m = re.search(r"(\d+) tests? collected", out)
        if m:
            ctx["tests"] = m.group(1)
    except Exception:
        pass
    ctx["cmp"] = pd.read_csv(ROOT / "reports" / "model_comparison.csv")
    ctx["meta"] = json.loads((ROOT / "models" / "model_meta.json").read_text())
    ctx["eda"] = json.loads((FIG / "eda_summary.json").read_text())
    return ctx


# --------------------------------------------------------------------------- helpers
def _text(
    slide,
    x,
    y,
    w,
    h,
    text,
    size=16,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    name=None,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT
    return box


def _bullets(slide, x, y, w, h, items, size=15, color=INK, gap=8):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, _, rest = item.partition("|")
        marker = p.add_run()
        marker.text = "▸  "
        marker.font.color.rgb = RED
        marker.font.size = Pt(size)
        marker.font.name = FONT
        r = p.add_run()
        r.text = head
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT
        r.font.bold = bool(rest)
        if rest:
            r2 = p.add_run()
            r2.text = " " + rest
            r2.font.size = Pt(size)
            r2.font.color.rgb = color
            r2.font.name = FONT
    return box


def _rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def _badge(slide, x, y, label, d=Inches(0.48), fill=RED, size=14):
    c = _rect(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = str(label)
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = FONT
    return c


def _image(slide, path: Path, x, y, w, h, caption=None, border=True):
    """Fit image into the box, preserving aspect; placeholder if the file is missing."""
    if path is None or not Path(path).exists():
        ph = _rect(slide, x, y, w, h, SAND, line=LINE)
        ph.line.dash_style = 4
        name = Path(path).name if path else "image"
        _text(
            slide,
            x,
            y,
            w,
            h,
            ["Add screenshot", f"screenshots/{name}"],
            size=13,
            color=MUTED,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        return
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph_ = int(iw * scale), int(ih * scale)
    px, py = x + (w - pw) // 2, y + (h - ph_) // 2
    pic = slide.shapes.add_picture(str(path), px, py, pw, ph_)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    if caption:
        _text(
            slide,
            x,
            y + h + Inches(0.05),
            w,
            Inches(0.3),
            caption,
            size=11,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )


SHOT_NAMES = {
    "01": "01_mlflow_experiments.png", "02": "02_mlflow_run_compare.png",
    "04": "04_github_actions_pipeline.png", "06": "06_github_actions_failure.png",
    "07": "07_docker_build_run.png", "08": "08_swagger_predict.png", "09": "09_k8s_resources.png",
    "10": "10_k8s_predict_ingress.png", "12": "12_grafana_dashboard.png", "13": "13_api_json_logs.png",
}


def _shot(n: str) -> Path:
    hits = sorted(SHOTS.glob(f"{n}_*.png"))
    return hits[0] if hits else SHOTS / SHOT_NAMES.get(n, f"{n}_screenshot.png")


def _content_slide(prs, title, kicker, number, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    _text(
        s,
        Inches(0.6),
        Inches(0.35),
        Inches(9),
        Inches(0.3),
        kicker.upper(),
        size=12,
        color=RED,
        bold=True,
    )
    _text(s, Inches(0.6), Inches(0.62), Inches(12), Inches(0.7), title, size=28, color=NAVY, bold=True)
    _rect(s, 0, H - Inches(0.42), W, Inches(0.42), SAND)
    _text(
        s,
        Inches(0.6),
        H - Inches(0.32),
        Inches(9),
        Inches(0.25),
        "Heart Disease Risk · End-to-End MLOps · AIMLCZG523 Assignment 01",
        size=10,
        color=MUTED,
    )
    _text(
        s,
        W - Inches(1.6),
        H - Inches(0.32),
        Inches(1.0),
        Inches(0.25),
        f"{number} / {total}",
        size=10,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )
    return s


def _stat(slide, x, y, w, value, label, color=NAVY):
    _text(slide, x, y, w, Inches(0.75), value, size=40, color=color, bold=True)
    _text(slide, x, y + Inches(0.78), w, Inches(0.5), label, size=12, color=MUTED)


# --------------------------------------------------------------------------- slides
def build():
    ctx = load_context()
    cmp_, meta, eda = ctx["cmp"], ctx["meta"], ctx["eda"]
    winner = meta["model_name"]
    wrow = cmp_.set_index("model").loc[winner]
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    TOTAL = 15
    n = iter(range(1, TOTAL + 1))

    # 1 ---- title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    next(n)
    _rect(s, Inches(0.6), Inches(1.55), Inches(0.12), Inches(2.1), RED)
    _text(
        s,
        Inches(0.95),
        Inches(1.45),
        Inches(11),
        Inches(0.5),
        "MLOPS · ASSIGNMENT 01",
        size=14,
        color=RGBColor(0xF4, 0xA2, 0xA8),
        bold=True,
    )
    _text(
        s,
        Inches(0.95),
        Inches(1.9),
        Inches(11.5),
        Inches(1.4),
        ["Heart Disease Risk Prediction", "End-to-End MLOps Pipeline"],
        size=40,
        color=WHITE,
        bold=True,
    )
    _text(
        s,
        Inches(0.95),
        Inches(3.45),
        Inches(11),
        Inches(0.5),
        "Data → Model → Tracking → CI/CD → Docker → Kubernetes → Monitoring",
        size=18,
        color=RGBColor(0xC9, 0xD3, 0xE6),
    )
    _text(
        s,
        Inches(0.95),
        Inches(5.3),
        Inches(11),
        Inches(1.2),
        [ctx["name"], "AIMLCZG523 Machine Learning Operations · BITS Pilani", ctx["repo"]],
        size=14,
        color=RGBColor(0xC9, 0xD3, 0xE6),
    )
    s.notes_slide.notes_text_frame.text = (
        "Introduce the goal: a heart-disease classifier operated like a production model — everything "
        "from data download to monitoring is automated and reproducible."
    )

    # 2 ---- what was delivered (9 tasks)
    s = _content_slide(prs, "Nine assignment tasks, one reproducible pipeline", "Scope", next(n), TOTAL)
    tasks = [
        ("Data & EDA", "download script, cleaning, 6 figures"),
        ("Features & models", f"sklearn Pipeline, {len(cmp_)} tuned models"),
        ("Experiment tracking", "MLflow runs, params, plots, registry"),
        ("Packaging", "model.joblib + meta, pinned deps"),
        ("CI/CD & tests", f"GitHub Actions, {ctx['tests']} pytest tests"),
        ("Containerisation", "FastAPI /predict in Docker"),
        ("Deployment", "Minikube: Deployment, LB, Ingress, HPA"),
        ("Monitoring", "JSON logs, Prometheus, Grafana"),
        ("Documentation", "report, README, architecture"),
    ]
    cw, ch = Inches(3.9), Inches(1.45)
    for i, (head, sub) in enumerate(tasks):
        cx = Inches(0.6) + (i % 3) * (cw + Inches(0.25))
        cy = Inches(1.6) + (i // 3) * (ch + Inches(0.2))
        _rect(s, cx, cy, cw, ch, SAND)
        _badge(s, cx + Inches(0.25), cy + Inches(0.28), i + 1)
        _text(
            s,
            cx + Inches(0.9),
            cy + Inches(0.25),
            cw - Inches(1.1),
            Inches(0.4),
            head,
            size=17,
            color=NAVY,
            bold=True,
        )
        _text(
            s,
            cx + Inches(0.9),
            cy + Inches(0.68),
            cw - Inches(1.1),
            Inches(0.6),
            sub,
            size=13,
            color=MUTED,
        )
    s.notes_slide.notes_text_frame.text = "Map each assignment task to what was built; details follow."

    # 3 ---- architecture
    s = _content_slide(
        prs, "Architecture: one package used by every stage", "Architecture", next(n), TOTAL
    )
    _image(
        s,
        ROOT / "docs" / "architecture.png",
        Inches(0.6),
        Inches(1.45),
        Inches(12.1),
        Inches(5.4),
        border=False,
    )
    s.notes_slide.notes_text_frame.text = (
        "src/heart is the single implementation of cleaning, features and inference; notebooks, "
        "training CLI, tests and the API all import it. CI builds the image that Kubernetes runs; "
        "Prometheus scrapes the pods."
    )

    # 4 ---- data
    s = _content_slide(
        prs, "Data: 303 Cleveland patients, nearly balanced", "Task 1 · Data acquisition", next(n), TOTAL
    )
    pos = eda["positive_rate"]
    miss = sum(eda.get("missing_raw", {}).values())
    _stat(s, Inches(0.6), Inches(1.6), Inches(2.6), str(eda["rows"]), "patients (UCI Cleveland)")
    _stat(s, Inches(3.2), Inches(1.6), Inches(2.6), f"{pos:.0%}", "have heart disease", RED)
    _stat(s, Inches(0.6), Inches(3.0), Inches(2.6), "13", "input features")
    _stat(s, Inches(3.2), Inches(3.0), Inches(2.6), str(miss), "missing values (ca, thal)", TEAL)
    _bullets(
        s,
        Inches(0.6),
        Inches(4.5),
        Inches(5.4),
        Inches(2.3),
        [
            "Download script|UCI first, GitHub mirror fallback",
            "Stateless cleaning|'?' → NaN, target num>0 → 1, range checks",
            "Imputation inside the pipeline|no leakage, same at serving",
        ],
        size=14,
    )
    _image(s, FIG / "01_class_balance.png", Inches(6.6), Inches(1.5), Inches(6.1), Inches(5.3))
    s.notes_slide.notes_text_frame.text = (
        "Demo: python scripts/download_data.py --force. Explain why imputation is not done in cleaning."
    )

    # 5 ---- EDA
    s = _content_slide(
        prs, "EDA: chest pain, thal and vessels carry the signal", "Task 1 · EDA", next(n), TOTAL
    )
    _image(
        s, FIG / "03_categorical_disease_rate.png", Inches(0.6), Inches(1.5), Inches(7.4), Inches(3.3)
    )
    _image(s, FIG / "04_correlation_heatmap.png", Inches(8.3), Inches(1.5), Inches(4.4), Inches(3.9))
    tops = list(eda["top_correlations_with_target"].items())[:4]
    _bullets(
        s,
        Inches(0.6),
        Inches(5.05),
        Inches(7.4),
        Inches(1.9),
        [
            "Strongest correlations with target:| " + ", ".join(f"{k} {v:.2f}" for k, v in tops),
            "Asymptomatic chest pain (cp=4) and reversible defect (thal=7)| → disease rate above 70 %",
            "chol, trestbps| overlap heavily between classes",
        ],
        size=13,
    )
    s.notes_slide.notes_text_frame.text = (
        "Mention the other figures in reports/figures (histograms, boxplots)."
    )

    # 6 ---- features pipeline
    s = _content_slide(
        prs, "Feature engineering lives inside the model pipeline", "Task 2 · Features", next(n), TOTAL
    )
    steps = [
        ("HeartFeatureBuilder", "derived features"),
        ("Impute", "median / mode"),
        ("Scale + One-hot", "ColumnTransformer"),
        ("Classifier", "tuned estimator"),
    ]
    bw, by = Inches(2.7), Inches(1.8)
    for i, (head, sub) in enumerate(steps):
        bx = Inches(0.6) + i * (bw + Inches(0.42))
        _rect(s, bx, by, bw, Inches(1.25), NAVY if i in (0, 3) else NAVY_2)
        _text(
            s,
            bx,
            by + Inches(0.22),
            bw,
            Inches(0.4),
            head,
            size=16,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        _text(
            s,
            bx,
            by + Inches(0.7),
            bw,
            Inches(0.4),
            sub,
            size=12,
            color=RGBColor(0xC9, 0xD3, 0xE6),
            align=PP_ALIGN.CENTER,
        )
        if i < 3:
            _text(
                s,
                bx + bw,
                by + Inches(0.38),
                Inches(0.42),
                Inches(0.5),
                "→",
                size=24,
                color=RED,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    _bullets(
        s,
        Inches(0.6),
        Inches(3.55),
        Inches(6.2),
        Inches(3.2),
        [
            "hr_reserve_pct|thalach / (220 − age)",
            "chol_per_age|cholesterol load relative to age",
            "st_risk_flag|oldpeak ≥ 2 or exercise angina",
            "One object saved|training and API apply identical transforms",
        ],
        size=14,
    )
    _image(
        s,
        FIG / "models" / f"{winner}_feature_importance.png",
        Inches(7.2),
        Inches(3.4),
        Inches(5.5),
        Inches(3.4),
    )
    s.notes_slide.notes_text_frame.text = (
        "Show src/heart/features.py. Point out the ablation model without engineered features."
    )

    # 7 ---- model comparison (native chart)
    s = _content_slide(
        prs, f"{winner} wins on cross-validated ROC-AUC", "Task 2 · Model selection", next(n), TOTAL
    )
    cd = CategoryChartData()
    cd.categories = list(cmp_["model"])
    cd.add_series("CV ROC-AUC (5-fold)", [round(v, 3) for v in cmp_["cv_roc_auc_mean"]])
    cd.add_series("Hold-out ROC-AUC", [round(v, 3) for v in cmp_["test_roc_auc"]])
    gf = s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.0), Inches(5.3), cd
    )
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    for ser, col in zip(ch.series, (NAVY, TEAL)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(11)
    plot.data_labels.number_format = "0.000"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = 0.8, 1.0
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(12)
    params = ", ".join(f"{k}={v}" for k, v in meta["best_params"].items())
    _rect(s, Inches(9.0), Inches(1.6), Inches(3.75), Inches(5.1), SAND)
    _text(
        s, Inches(9.3), Inches(1.85), Inches(3.2), Inches(0.3), "SELECTED", size=12, color=RED, bold=True
    )
    _text(s, Inches(9.3), Inches(2.15), Inches(3.2), Inches(0.5), winner, size=24, color=NAVY, bold=True)
    _text(s, Inches(9.3), Inches(2.7), Inches(3.2), Inches(0.6), params, size=12, color=MUTED)
    _stat(
        s,
        Inches(9.3),
        Inches(3.45),
        Inches(3.2),
        f"{wrow.cv_roc_auc_mean:.3f}",
        f"CV ROC-AUC  (± {wrow.cv_roc_auc_std:.3f})",
    )
    _stat(s, Inches(9.3), Inches(4.8), Inches(3.2), f"{wrow.test_recall:.0%}", "hold-out recall", RED)
    s.notes_slide.notes_text_frame.text = (
        "Rule: highest mean 5-fold CV ROC-AUC on the training split; hold-out used only for reporting. "
        "Linear models beat tree ensembles on 303 rows."
    )

    # 8 ---- evaluation
    s = _content_slide(prs, "Selected model on the hold-out set", "Task 2 · Evaluation", next(n), TOTAL)
    _image(
        s,
        FIG / "models" / f"{winner}_evaluation.png",
        Inches(0.6),
        Inches(1.5),
        Inches(12.1),
        Inches(3.9),
    )
    stats = [
        (f"{wrow.test_accuracy:.1%}", "accuracy"),
        (f"{wrow.test_precision:.1%}", "precision"),
        (f"{wrow.test_recall:.1%}", "recall"),
        (f"{wrow.test_roc_auc:.3f}", "ROC-AUC"),
    ]
    for i, (v, lab) in enumerate(stats):
        _stat(
            s,
            Inches(0.9) + i * Inches(3.1),
            Inches(5.55),
            Inches(2.8),
            v,
            lab,
            RED if lab == "recall" else NAVY,
        )
    s.notes_slide.notes_text_frame.text = (
        "Recall matters most clinically: missing a disease case is costly."
    )

    # 9 ---- MLflow
    s = _content_slide(
        prs, "Every run tracked in MLflow", "Task 3 · Experiment tracking", next(n), TOTAL
    )
    _image(
        s,
        _shot("01"),
        Inches(0.6),
        Inches(1.5),
        Inches(6.0),
        Inches(3.6),
        caption="Parent + nested runs",
    )
    _image(s, _shot("02"), Inches(6.75), Inches(1.5), Inches(6.0), Inches(3.6), caption="Run comparison")
    _bullets(
        s,
        Inches(0.6),
        Inches(5.55),
        Inches(12),
        Inches(1.4),
        [
            "Params|best hyper-parameters, grid size, feature set     ",
            "Metrics|CV mean ± std for 5 metrics, hold-out metrics     ",
            "Artifacts|plots, grid results, model with signature, registry entry",
        ],
        size=13,
        gap=4,
    )
    s.notes_slide.notes_text_frame.text = "Demo: make mlflow → http://localhost:5000, open a child run."

    # 10 ---- packaging & tests
    s = _content_slide(
        prs,
        "Reproducible packaging and automated tests",
        "Tasks 4–5 · Packaging & tests",
        next(n),
        TOTAL,
    )
    _bullets(
        s,
        Inches(0.6),
        Inches(1.6),
        Inches(6.0),
        Inches(5),
        [
            "models/model.joblib|whole pipeline in one object",
            "model_meta.json|version, run IDs, params, metrics, feature order",
            "3-layer requirements|serve · train · dev, all pinned",
            "Fixed seeds|RANDOM_STATE = 42 everywhere",
            f"{ctx['tests']} pytest tests|data, features, model, API — no network needed",
            "Make targets|make data · train · test · lint · docker-build",
        ],
        size=15,
        gap=10,
    )
    _rect(s, Inches(7.0), Inches(1.6), Inches(5.75), Inches(5.1), NAVY)
    _text(
        s,
        Inches(7.3),
        Inches(1.85),
        Inches(5.2),
        Inches(4.6),
        [
            "$ make lint",
            "python -m flake8",
            "",
            "$ make test",
            f"{ctx['tests']} passed",
            "",
            "$ make train",
            f"Selected {winner} -> models/model.joblib",
            "",
            "$ make docker-build docker-run smoke",
            "SMOKE TEST PASSED",
            "",
            "$ ./scripts/deploy_minikube.sh",
            "deployment \"heart-api\" successfully rolled out",
        ],
        size=15,
        color=RGBColor(0xC9, 0xD3, 0xE6),
    )
    s.notes_slide.notes_text_frame.text = "Demo: run make lint test live in the terminal."

    # 11 ---- CI/CD
    s = _content_slide(prs, "GitHub Actions gates every change", "Task 5 · CI/CD", next(n), TOTAL)
    jobs = ["lint", "k8s-validate", "test", "train + gate", "docker"]
    jw = Inches(2.25)
    for i, j in enumerate(jobs):
        jx = Inches(0.6) + i * (jw + Inches(0.22))
        _rect(
            s,
            jx,
            Inches(1.55),
            jw,
            Inches(0.7),
            TEAL if i < 4 else NAVY,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        )
        _text(
            s,
            jx,
            Inches(1.55),
            jw,
            Inches(0.7),
            j,
            size=15,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    _image(
        s, _shot("04"), Inches(0.6), Inches(2.55), Inches(7.6), Inches(3.75), caption="Green run on main"
    )
    _image(
        s,
        _shot("06"),
        Inches(8.45),
        Inches(2.55),
        Inches(4.3),
        Inches(3.75),
        caption="Lint error stops the pipeline",
    )
    s.notes_slide.notes_text_frame.text = (
        "Quality gate: CV ROC-AUC ≥ 0.85. Artifacts: lint report, JUnit, coverage, model, MLflow store, "
        "container logs. Image pushed to GHCR on main."
    )

    # 12 ---- container & API
    s = _content_slide(
        prs, "FastAPI service in a slim, non-root container", "Task 6 · Containerisation", next(n), TOTAL
    )
    rows = [
        ("POST /predict", "prediction, confidence, risk band"),
        ("POST /predict/batch", "up to 500 records"),
        ("GET /health", "200 when model loaded"),
        ("GET /metrics", "Prometheus format"),
        ("GET /docs", "Swagger UI"),
    ]
    tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.6), Inches(1.6), Inches(5.6), Inches(3.0)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(2.3), Inches(3.3)
    for c, head in enumerate(("Endpoint", "Returns")):
        cell = tbl.cell(0, c)
        cell.text = head
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for r, (a, b) in enumerate(rows, start=1):
        for c, val in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = SAND if r % 2 else WHITE
    for r in range(len(rows) + 1):
        for c in range(2):
            for p in tbl.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(13)
                    run.font.name = FONT
                    run.font.color.rgb = WHITE if r == 0 else INK
                    run.font.bold = r == 0 or c == 0
    _bullets(
        s,
        Inches(0.6),
        Inches(4.95),
        Inches(5.6),
        Inches(1.9),
        [
            "python:3.11-slim|serving deps only, UID 1001, HEALTHCHECK",
            "Pydantic validation|bad input → 422",
        ],
        size=13,
        gap=6,
    )
    _image(
        s, _shot("08"), Inches(6.5), Inches(1.5), Inches(6.25), Inches(2.55), caption="Swagger /predict"
    )
    _image(
        s, _shot("07"), Inches(6.5), Inches(4.3), Inches(6.25), Inches(2.3), caption="Build + smoke test"
    )
    s.notes_slide.notes_text_frame.text = "Demo: make docker-build docker-run smoke."

    # 13 ---- Kubernetes
    s = _content_slide(prs, "Deployed on Kubernetes (Minikube)", "Task 7 · Deployment", next(n), TOTAL)
    _bullets(
        s,
        Inches(0.6),
        Inches(1.6),
        Inches(4.6),
        Inches(5.2),
        [
            "Deployment|2 replicas, rolling update, probes, limits",
            "Service|LoadBalancer on :8080 (minikube tunnel)",
            "Ingress|nginx, host heart-api.local",
            "HPA|2–5 replicas at 70 % CPU",
            "Kustomize|kubectl apply -k k8s/",
            "One script|./scripts/deploy_minikube.sh",
        ],
        size=14,
        gap=10,
    )
    _image(
        s, _shot("09"), Inches(5.5), Inches(1.5), Inches(7.25), Inches(2.45), caption="Cluster resources"
    )
    _image(
        s,
        _shot("10"),
        Inches(5.5),
        Inches(4.3),
        Inches(7.25),
        Inches(2.3),
        caption="Prediction through the Ingress",
    )
    s.notes_slide.notes_text_frame.text = (
        "Demo: kubectl -n heart-mlops get pods,svc,ingress,hpa; curl http://heart-api.local/health"
    )

    # 14 ---- monitoring
    s = _content_slide(
        prs, "Monitoring: logs, metrics, dashboard", "Task 8 · Monitoring", next(n), TOTAL
    )
    _image(
        s, _shot("12"), Inches(0.6), Inches(1.5), Inches(7.7), Inches(4.3), caption="Grafana dashboard"
    )
    _image(
        s, _shot("13"), Inches(8.55), Inches(1.5), Inches(4.2), Inches(2.0), caption="JSON request logs"
    )
    _bullets(
        s,
        Inches(8.55),
        Inches(4.0),
        Inches(4.2),
        Inches(2.8),
        [
            "Request rate, p95 latency, errors",
            "Predictions per class",
            "Mean P(disease)|drift signal",
            "Prometheus pod discovery",
        ],
        size=13,
        gap=6,
    )
    s.notes_slide.notes_text_frame.text = (
        "Demo: python scripts/load_test.py --url http://heart-api.local, "
        "then show Grafana panels updating."
    )

    # 15 ---- close
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    next(n)
    _text(
        s,
        Inches(0.8),
        Inches(0.8),
        Inches(11),
        Inches(0.5),
        "SUMMARY",
        size=14,
        color=RGBColor(0xF4, 0xA2, 0xA8),
        bold=True,
    )
    _text(
        s,
        Inches(0.8),
        Inches(1.25),
        Inches(11.5),
        Inches(0.9),
        "From raw CSV to a monitored API — fully automated",
        size=32,
        color=WHITE,
        bold=True,
    )
    summary = [
        (f"{wrow.cv_roc_auc_mean:.3f}", "CV ROC-AUC"),
        (f"{wrow.test_recall:.0%}", "hold-out recall"),
        (ctx["tests"], "unit tests"),
        ("5", "CI jobs"),
    ]
    for i, (v, lab) in enumerate(summary):
        x = Inches(0.8) + i * Inches(3.0)
        _text(s, x, Inches(2.6), Inches(2.8), Inches(0.9), v, size=44, color=WHITE, bold=True)
        _text(
            s, x, Inches(3.5), Inches(2.8), Inches(0.4), lab, size=14, color=RGBColor(0xC9, 0xD3, 0xE6)
        )
    _text(
        s,
        Inches(0.8),
        Inches(4.5),
        Inches(11.5),
        Inches(0.4),
        "Next steps",
        size=16,
        color=RGBColor(0xF4, 0xA2, 0xA8),
        bold=True,
    )
    _text(
        s,
        Inches(0.8),
        Inches(4.9),
        Inches(11.5),
        Inches(1.2),
        [
            "Probability calibration and a recall-oriented threshold · "
            "per-feature drift tests on logged inputs",
            "Remote MLflow server with model promotion · Helm chart with separate environments",
        ],
        size=15,
        color=RGBColor(0xC9, 0xD3, 0xE6),
    )
    _text(s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4), ctx["repo"], size=14, color=WHITE)
    s.notes_slide.notes_text_frame.text = "Thank the audience, share the repo link, take questions."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    missing = [
        k
        for k in ("01", "02", "04", "06", "07", "08", "09", "10", "12", "13")
        if not list(SHOTS.glob(f"{k}_*.png"))
    ]
    print(f"wrote {OUT}")
    if missing:
        print("screenshots missing (placeholders used):", ", ".join(missing))


if __name__ == "__main__":
    build()
