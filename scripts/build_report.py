"""Render reports/REPORT.md into reports/Heart_Disease_MLOps_Report.docx.

Re-run after adding screenshots: missing images become visible placeholders.

    python scripts/build_report.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "reports" / "REPORT.md"
OUT = ROOT / "reports" / "Heart_Disease_MLOps_Report.docx"
ACCENT = RGBColor(0x26, 0x46, 0x53)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|<https?://[^>]+>|<[^>]+>)")


def add_inline(par, text: str, size: float | None = None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x9b, 0x2c, 0x2c)
            continue
        elif part.startswith("<"):
            run = par.add_run(part[1:-1])
            run.font.color.rgb = RGBColor(0x1d, 0x4e, 0x89)
        else:
            run = par.add_run(part)
        if size:
            run.font.size = Pt(size)


def shade(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, rows: list[list[str]]):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8 if len(header) > 5 else 9
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], h, size)
        for r in cell.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        shade(cell, "264653")
    for k, row in enumerate(body):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], val, size)
            if k % 2:
                shade(cells[i], "EEF4F3")
    doc.add_paragraph()


def add_image(doc, alt: str, rel: str):
    path = (SRC.parent / rel).resolve()
    if path.exists():
        from PIL import Image
        w, h = Image.open(path).size
        width = 16 if w / h > 1.8 else (12 if w / h > 1.2 else 9.5)
        doc.add_picture(str(path), width=Cm(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"[ Insert screenshot: {path.relative_to(ROOT)} ]")
        run.italic = True
        run.font.color.rgb = RGBColor(0xb0, 0x30, 0x30)
        p_pr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "dashed")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:space"), "12")
            el.set(qn("w:color"), "B03030")
            border.append(el)
        p_pr.append(border)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(alt)
    r.italic = True
    r.font.size = Pt(9)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.color.rgb = ACCENT
        st.font.size = Pt(size)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AIMLCZG523 · Assignment 01 · Heart Disease MLOps").font.size = Pt(8)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            if first_h1:
                t = doc.add_paragraph()
                t.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = t.add_run(line[2:])
                run.bold = True
                run.font.size = Pt(20)
                run.font.color.rgb = ACCENT
                first_h1 = False
            else:
                doc.add_heading(line[2:], 1)
        elif line.startswith("## "):
            if line.startswith("## 1."):
                doc.add_page_break()
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            add_image(doc, m.group(1), m.group(2))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^(\s*)[-*] ", line) or re.match(r"^\d+\. ", line):
            numbered = bool(re.match(r"^\d+\. ", line))
            text = re.sub(r"^(\s*)([-*]|\d+\.) ", "", line)
            i += 1
            while (i < len(lines) and lines[i].startswith("  ") and lines[i].strip()
                   and not re.match(r"^\s*([-*]|\d+\.) ", lines[i])):
                text += " " + lines[i].strip()
                i += 1
            add_inline(doc.add_paragraph(style="List Number" if numbered else "List Bullet"), text)
            continue
        else:
            text = line
            i += 1
            block_start = re.compile(r"^(#|!\[|\||[-*] |\d+\. )")
            while i < len(lines) and lines[i].strip() and not block_start.match(lines[i]):
                sep = "\n" if text.endswith("  ") or re.match(r"^\*\*[^*]+:\*\*", lines[i - 1]) else " "
                text += sep + lines[i].strip()
                i += 1
            for k, chunk in enumerate(text.split("\n")):
                p = doc.add_paragraph()
                if first_h1 is False and doc.paragraphs and len(doc.paragraphs) < 6:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, chunk.strip())
            continue
        i += 1
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
